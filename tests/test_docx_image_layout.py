from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from geng_agent.docx_assets import (
    _add_image_comparison_table,
    _add_markdown_image,
    _page_content_size,
)
from geng_agent.docx_styles import _setup_document


class DocxImageLayoutTests(unittest.TestCase):
    def _image(self, root: Path, name: str, size: tuple[int, int], *, dpi=(96, 96)) -> Path:
        path = root / name
        Image.new("RGB", size, color="white").save(path, dpi=dpi)
        return path

    def _document(self):
        document = Document()
        _setup_document(document)
        return document

    def _text(self, document) -> str:
        return "\n".join(
            paragraph.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        )

    def test_similarly_shaped_single_figures_remain_side_by_side(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            local = self._image(root, "local.png", (900, 750))
            paper = self._image(root, "paper.png", (1000, 800))
            document = self._document()
            _add_image_comparison_table(document, ["本地图", "论文图"], [[f"![本地说明]({local.name})", f"![图 1，第 3 页]({paper.name})"]], base_dir=root)

            table = document.tables[0]
            self.assertEqual(len(table.rows), 2)
            self.assertEqual(len(table.rows[1]._tr.tc_lst), 2)
            self.assertEqual(len(document.inline_shapes), 2)
            for shape, ratio in zip(document.inline_shapes, (1.2, 1.25)):
                self.assertLess(shape.width.inches, _page_content_size(document)[0] / 2)
                self.assertAlmostEqual(shape.width / shape.height, ratio, places=5)
            self.assertIn("图 1，第 3 页", self._text(document))
            self.assertEqual(document.inline_shapes[1]._inline.docPr.get("title"), paper.name)
            self.assertEqual(document.inline_shapes[1]._inline.docPr.get("descr"), "图 1，第 3 页")

    def test_wide_plot_and_full_paper_page_are_stacked_without_cropping(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            wide = self._image(root, "wide.png", (1600, 600))
            page = self._image(root, "page.png", (1000, 1414))
            document = self._document()
            _add_image_comparison_table(document, ["本地组合图", "论文原页"], [[f"![多个子图]({wide})", f"![论文第 2 页]({page})"]])

            self.assertEqual(len(document.tables[0].rows), 4)
            self.assertTrue(all(len(row._tr.tc_lst) == 1 for row in document.tables[0].rows))
            self.assertEqual(len(document.inline_shapes), 2)
            width, height = _page_content_size(document)
            for shape, ratio in zip(document.inline_shapes, (1600 / 600, 1000 / 1414)):
                self.assertGreater(shape.width.inches, 5)
                self.assertLessEqual(shape.width.inches, width)
                self.assertLess(shape.height.inches, height - 0.65)
                self.assertAlmostEqual(shape.width / shape.height, ratio, places=5)
                self.assertFalse(shape._inline.xpath(".//a:srcRect"))
            self.assertIn("论文第 2 页", self._text(document))

    def test_very_tall_standalone_image_fits_short_page_and_preserves_pixel_ratio(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            image = self._image(root, "tall.png", (500, 2400), dpi=(72, 144))
            document = self._document()
            section = document.sections[0]
            section.page_height = Inches(7)
            section.page_width = Inches(5)
            _add_markdown_image(document, "完整长图", str(image))

            shape = document.inline_shapes[0]
            width, height = _page_content_size(document)
            self.assertLessEqual(shape.width.inches, width)
            self.assertLess(shape.height.inches, height - 0.65)
            self.assertAlmostEqual(shape.width / shape.height, 500 / 2400, places=5)
            self.assertTrue(document.paragraphs[0].paragraph_format.keep_with_next)
            self.assertFalse(document.paragraphs[1].paragraph_format.keep_with_next)
            self.assertEqual(document.paragraphs[1].text, "完整长图")

    def test_multifigure_missing_and_text_cells_keep_all_content(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            first = self._image(root, "first.png", (800, 600))
            second = self._image(root, "second.png", (800, 600))
            document = self._document()
            _add_image_comparison_table(document, ["本地证据", "原文证据", "说明"], [[f"![曲线 A]({first.name})<br />![曲线 B]({second.name})", "![原文图 3，第 7 页](missing.png)", "图例中的并列关系需要人工核查。"]], base_dir=root)

            self.assertEqual(len(document.inline_shapes), 2)
            for shape in document.inline_shapes:
                self.assertGreater(shape.width.inches, 5)
            text = self._text(document)
            for fragment in ("本地证据", "曲线 A", "曲线 B", "原文图 3，第 7 页", "图片缺失：missing.png", "图例中的并列关系需要人工核查。"):
                self.assertIn(fragment, text)
            self.assertNotIn("![", text)
            for row in document.tables[0].rows:
                self.assertLessEqual(len(row._tr.xpath(".//w:drawing")), 1)
                # Multiple full-width images are not locked into one giant row.
                self.assertFalse(row._tr.xpath("./w:trPr/w:cantSplit"))

    def test_corrupt_image_keeps_original_caption_and_delivery_warning(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "broken.png").write_bytes(b"not a png")
            document = self._document()
            _add_image_comparison_table(document, ["论文图"], [["![图 4 原文第 9 页](broken.png)"]], base_dir=root)
            self.assertEqual(len(document.inline_shapes), 0)
            self.assertIn("图 4 原文第 9 页", self._text(document))
            self.assertIn("图片插入失败：broken.png", self._text(document))

    def test_long_alt_and_caption_are_not_truncated(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            image = self._image(root, "figure.png", (1000, 600))
            caption = "原文第 3 页图 2，条件与图例说明。" * 100 + "完整末尾"
            document = self._document()
            _add_markdown_image(document, caption, str(image))
            self.assertEqual(document.paragraphs[1].text, caption)
            self.assertEqual(document.inline_shapes[0]._inline.docPr.get("descr"), caption)
            self.assertNotIn("已截断", document.paragraphs[1].text)
            self.assertFalse(document.paragraphs[1].paragraph_format.keep_together)

    def test_dissimilar_shapes_stack_even_when_neither_is_extremely_wide(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            square = self._image(root, "square.png", (800, 800))
            landscape = self._image(root, "landscape.png", (1200, 800))
            document = self._document()
            _add_image_comparison_table(document, ["A", "B"], [[f"![A]({square})", f"![B]({landscape})"]])
            self.assertTrue(all(len(row._tr.tc_lst) == 1 for row in document.tables[0].rows))
            self.assertGreater(document.inline_shapes[0].width.inches, 5)

    def test_comparison_width_tracks_custom_section_and_survives_save(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            image = self._image(root, "figure.png", (1200, 600))
            document = self._document()
            section = document.sections[0]
            section.page_width = Inches(6)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
            _add_image_comparison_table(document, ["A", "B"], [[f"![A]({image})", f"![B]({image})"]])
            path = root / "figures.docx"
            document.save(path)
            reopened = Document(path)
            self.assertEqual(len(reopened.inline_shapes), 2)
            self.assertEqual(reopened.tables[0]._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")), "5760")
            for shape in reopened.inline_shapes:
                self.assertLessEqual(shape.width.inches, 4)
                self.assertGreater(shape.width.inches, 3.5)


if __name__ == "__main__":
    unittest.main()
