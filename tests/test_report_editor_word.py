from __future__ import annotations

from docx import Document

from geng_agent.pipeline_report_delivery import inspect_editor_word_reports
from geng_agent.report_editor_word import inspect_word_file


def test_editor_authored_word_is_delivered_byte_for_byte(tmp_path):
    for stem in ("reproduction_report", "result_review"):
        report = Document()
        report.add_heading("任务 T1", 1)
        report.add_paragraph("智能体写出的中文结论。")
        report.save(tmp_path / f"{stem}.docx")
    original = (tmp_path / "result_review.docx").read_bytes()

    status = inspect_editor_word_reports(output_dir=tmp_path, result_review_result={"passed": True})

    assert status["review_docx"]["passed"] is None
    assert status["result_review_docx"]["passed"] is True
    assert status["reproduction_report_docx"]["passed"] is True
    assert (tmp_path / "result_review.docx").read_bytes() == original


def test_damaged_word_is_reported_without_generating_a_replacement(tmp_path):
    (tmp_path / "result_review.docx").write_bytes(b"not a Word document")

    status = inspect_editor_word_reports(output_dir=tmp_path, result_review_result={"passed": True})

    assert status["result_review_docx"]["passed"] is False
    assert status["reproduction_report_docx"]["passed"] is False
    assert inspect_word_file(tmp_path / "result_review.docx")
    assert (tmp_path / "result_review.docx").read_bytes() == b"not a Word document"
    assert not (tmp_path / "reproduction_report.docx").exists()
